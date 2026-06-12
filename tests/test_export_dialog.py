import csv
import os
import tempfile

import pytest
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QApplication

from pdf_gui.models.config import (FlowConfig, JobColumnConfig, MetricConfig,
                                   StepConfig, StepGroupConfig)
from pdf_gui.models.run_data import (Job, OverallStatus, Step, StepStatus,
                                     Version)
from pdf_gui.widgets.export_dialog import ExportDialog


def make_config():
    return FlowConfig(
        flow_name="Test",
        step_groups=[StepGroupConfig(
            name="All",
            steps=["init", "place"],
            metrics=[
                MetricConfig(key="WNS", format=".3f"),
                MetricConfig(key="TNS", format=".2f"),
            ],
        )],
        job_columns=[
            JobColumnConfig(key="status"),
            JobColumnConfig(key="runtime"),
        ],
    )


def make_version(name: str, status: OverallStatus, wns=None, tns=None,
                 step_status=StepStatus.SUCCESS):
    steps = []
    for step_name in ["init", "place"]:
        s = Step(
            name=step_name,
            status=step_status,
            metrics={"WNS": wns, "TNS": tns},
            job=Job(job_id="12345", status=step_status, runtime="10m"),
        )
        steps.append(s)
    return Version(name=name, dir_path="/tmp", status=status,
                   steps=steps, latest_step=steps[-1].name)


def test_export_dialog_all_checked():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])

    versions = [
        make_version("v1", OverallStatus.SUCCESS),
        make_version("v2", OverallStatus.RUNNING),
        make_version("v3", OverallStatus.FAIL),
    ]
    dialog = ExportDialog(versions, make_config())
    for i in range(dialog._list.count()):
        assert dialog._list.item(i).checkState() == Qt.Checked


def test_csv_export_selected_versions():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])

    config = make_config()
    versions = [
        make_version("v1", OverallStatus.SUCCESS, wns=-0.050, tns=-3.45),
        make_version("v2", OverallStatus.RUNNING, wns=-0.100, tns=-8.00),
        make_version("v3", OverallStatus.FAIL, wns=-0.500, tns=-20.00),
    ]

    dialog = ExportDialog(versions, config)
    # select only v1 and v3
    dialog._list.item(0).setCheckState(Qt.Checked)
    dialog._list.item(1).setCheckState(Qt.Unchecked)
    dialog._list.item(2).setCheckState(Qt.Checked)

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "test.csv")
        dialog._path_edit.setText(path)
        dialog._write_csv(path, dialog._selected_versions())

        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = list(reader)

        assert rows[0] == ["version: v1"]
        assert rows[1] == ["# All"]
        assert rows[2] == ["Step", "WNS", "TNS", "Status", "Runtime"]
        assert rows[3][0] == "Init"
        assert rows[3][1] == "-0.050"
        assert rows[3][2] == "-3.45"
        assert rows[3][3] == "✓ SUCCESS"
        assert rows[3][4] == "10m"
        # v2 should NOT appear
        version_headers = [r[0] for r in rows if r and r[0].startswith("version:")]
        assert len(version_headers) == 2
        assert "v1" in version_headers[0]
        assert "v3" in version_headers[1]


def test_csv_export_empty_selection():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])

    config = make_config()
    versions = [
        make_version("v1", OverallStatus.SUCCESS),
    ]
    dialog = ExportDialog(versions, config)
    dialog._list.item(0).setCheckState(Qt.Unchecked)

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "empty.csv")
        dialog._path_edit.setText(path)
        # should not raise, writes empty file with no version sections
        dialog._write_csv(path, dialog._selected_versions())
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            content = f.read()
        assert "version:" not in content


def test_format_switches_extension():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])

    dialog = ExportDialog([], make_config())
    dialog._path_edit.setText("export.xlsx")

    dialog._format_combo.setCurrentIndex(0)  # switch to CSV
    assert dialog._path_edit.text() == "export.csv"

    dialog._format_combo.setCurrentIndex(1)  # back to Excel
    assert dialog._path_edit.text() == "export.xlsx"


def test_csv_per_group_format():
    """CSV has per-version group sections, only groups with steps."""
    app = QApplication.instance()
    if app is None:
        app = QApplication([])

    from pdf_gui.models.config import StepGroupConfig
    from pdf_gui.models.run_data import OverallStatus, Step, StepStatus, Version

    config = FlowConfig(
        step_groups=[
            StepGroupConfig(name="APR", steps=["init", "place"],
                            metrics=[MetricConfig(key="WNS", format=".3f")]),
            StepGroupConfig(name="STA", steps=["sta"],
                            metrics=[MetricConfig(key="WNS", format=".3f")]),
        ],
        job_columns=[JobColumnConfig(key="status")],
    )
    versions = [
        Version(name="v1", dir_path="/tmp/a", status=OverallStatus.SUCCESS,
                steps=[
                    Step(name="init", status=StepStatus.SUCCESS,
                         metrics={"WNS": -0.050},
                         job=Job(job_id="1", status=StepStatus.SUCCESS)),
                    Step(name="place", status=StepStatus.SUCCESS,
                         metrics={"WNS": -0.100},
                         job=Job(job_id="1", status=StepStatus.SUCCESS)),
                ], latest_step="place"),
    ]
    dialog = ExportDialog(versions, config)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "test.csv")
        dialog._write_csv(path, versions)
        with open(path, "r", encoding="utf-8-sig") as f:
            content = f.read()
        assert "version: v1" in content
        assert "# APR" in content
        assert "# STA" not in content  # v1 has no STA steps


def test_version_without_group_omitted():
    """Version without group steps not in XLSX sheet."""
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    pytest.importorskip("openpyxl")

    from pdf_gui.models.config import StepGroupConfig

    config = FlowConfig(
        step_groups=[
            StepGroupConfig(name="APR", steps=["init"],
                            metrics=[MetricConfig(key="WNS", format=".3f")]),
            StepGroupConfig(name="STA", steps=["sta"],
                            metrics=[MetricConfig(key="WNS", format=".3f")]),
        ],
        job_columns=[JobColumnConfig(key="status")],
    )
    versions = [
        Version(name="v_apr_only", dir_path="/tmp/a",
                status=OverallStatus.SUCCESS,
                steps=[Step(name="init", status=StepStatus.SUCCESS,
                            metrics={"WNS": -0.050})],
                latest_step="init"),
    ]
    dialog = ExportDialog(versions, config)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "test.xlsx")
        dialog._write_excel(path, versions)
        import openpyxl
        wb = openpyxl.load_workbook(path)
        assert "APR" in wb.sheetnames
        assert "STA" in wb.sheetnames
        # STA sheet should be empty (v_apr_only has no sta step)
        ws_sta = wb["STA"]
        assert ws_sta.cell(1, 1).value is None  # no header = no data


def test_xlsx_export():
    pytest.importorskip("openpyxl")
    import openpyxl

    app = QApplication.instance()
    if app is None:
        app = QApplication([])

    config = make_config()
    versions = [
        make_version("v1", OverallStatus.SUCCESS, wns=-0.050, tns=-3.45),
        make_version("v2", OverallStatus.RUNNING, wns=-0.100, tns=-8.00),
    ]
    dialog = ExportDialog(versions, config)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "test.xlsx")
        dialog._write_excel(path, versions)
        assert os.path.isfile(path)
        wb = openpyxl.load_workbook(path)
        assert "Versions" in wb.sheetnames or "All" in wb.sheetnames
        ws = wb.active

        # header row
        assert ws.cell(1, 1).value == "Version"
        assert ws.cell(1, 2).value == "Step"

        # v1: cell A2 = "v1\n(SUCCESS)", merged across rows
        assert ws.cell(2, 1).value == "v1\n(SUCCESS)"
        assert ws.cell(2, 2).value == "Init"

        # verify merge
        merged = [str(m) for m in ws.merged_cells.ranges]
        assert any("A2" in m for m in merged)


def test_csv_grouped_header():
    """CSV export includes @-prefixed metric key labels in header."""
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    from pdf_gui.models.config import StepGroupConfig
    config = FlowConfig(
        step_groups=[StepGroupConfig(
            name="All", steps=["init"],
            metrics=[
                MetricConfig(key="density", label="Density"),
                MetricConfig(key="REG2REG@wns", label="wns"),
                MetricConfig(key="REG2REG@tns", label="tns"),
            ],
        )],
        job_columns=[JobColumnConfig(key="status")],
    )
    versions = [
        Version(name="v1", dir_path="/tmp/a", status=OverallStatus.SUCCESS,
                steps=[Step(name="init", status=StepStatus.SUCCESS,
                            metrics={"density": 56.0, "REG2REG@wns": -0.1,
                                     "REG2REG@tns": -3.0})],
                latest_step="init"),
    ]
    dialog = ExportDialog(versions, config)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "test.csv")
        dialog._write_csv(path, versions)
        with open(path, "r", encoding="utf-8-sig") as f:
            content = f.read()
        # grouped header: REG2REG parent label in row0, wns/tns sub-labels in row1
        assert "REG2REG" in content
        assert "wns" in content
        assert "Density" in content
